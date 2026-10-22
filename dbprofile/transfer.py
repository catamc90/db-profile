"""Import / export engine (no UI): file formats, SQL dumps and import modes.

Everything here is blocking and meant to run in a worker thread. Long operations
report progress through a callback ``progress(done, total, message)`` and stop
early when ``cancelled()`` returns True.
"""
import base64
import csv
import datetime
import decimal
import html
import json
import os
import re
from xml.sax.saxutils import escape as xml_escape

import pymysql

from . import APP_NAME, APP_VERSION, db

try:
    import openpyxl
except ImportError:  # optional dependency
    openpyxl = None

q = db.quote_ident

EXPORT_FORMATS = {  # key: (label, file extension)
    "csv": ("CSV file (*.csv)", "csv"),
    "txt": ("Text file (*.txt)", "txt"),
    "json": ("JSON file (*.json)", "json"),
    "xml": ("XML file (*.xml)", "xml"),
    "html": ("HTML file (*.html)", "html"),
    "sql": ("SQL script file (*.sql)", "sql"),
}
if openpyxl:
    EXPORT_FORMATS["xlsx"] = ("Excel file (*.xlsx)", "xlsx")

IMPORT_FORMATS = {"csv": "CSV file (*.csv)", "txt": "Text file (*.txt)", "json": "JSON file (*.json)"}
if openpyxl:
    IMPORT_FORMATS["xlsx"] = "Excel file (*.xlsx)"

IMPORT_MODES = {
    "append": "Append: add records to the destination table",
    "update": "Update: update records in the destination with matching source records",
    "append_update": "Append/Update: update matching records, add the others",
    "delete": "Delete: delete records in the destination that match the source records",
    "copy": "Copy: delete all records in the destination, then append the source records",
}

CHUNK = 2000


class Cancelled(Exception):
    pass


def _never():
    return False


def _noop(*_):
    pass


# --- value conversion ------------------------------------------------------------
def to_text(value, null_text=""):
    if value is None:
        return null_text
    if isinstance(value, (bytes, bytearray)):
        return base64.b64encode(bytes(value)).decode("ascii")
    if isinstance(value, decimal.Decimal):
        return format(value, "f")
    return str(value)


def to_json(value):
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, decimal.Decimal):
        return format(value, "f")
    if isinstance(value, (bytes, bytearray)):
        return base64.b64encode(bytes(value)).decode("ascii")
    return str(value)


# --- writers -----------------------------------------------------------------------
class _Writer:
    def __init__(self, path, columns, options, name):
        self.path, self.columns, self.options, self.name = path, columns, options, name

    def write_rows(self, rows):
        raise NotImplementedError

    def close(self):
        pass


class DelimitedWriter(_Writer):
    def __init__(self, path, columns, options, name):
        super().__init__(path, columns, options, name)
        self.f = open(path, "w", newline="", encoding=options.get("encoding", "utf-8"))
        quote = options.get("quote", '"')
        self.w = csv.writer(
            self.f,
            delimiter=options.get("delimiter", ","),
            quotechar=quote or '"',
            quoting=csv.QUOTE_MINIMAL if quote else csv.QUOTE_NONE,
            escapechar=None if quote else "\\",
            lineterminator=options.get("line_ending", "\r\n"),
        )
        self.null = options.get("null_text", "")
        if options.get("header", True):
            self.w.writerow(columns)

    def write_rows(self, rows):
        self.w.writerows([to_text(v, self.null) for v in row] for row in rows)

    def close(self):
        self.f.close()


class JsonWriter(_Writer):
    def __init__(self, path, columns, options, name):
        super().__init__(path, columns, options, name)
        self.f = open(path, "w", encoding="utf-8")
        self.indent = 2 if options.get("pretty", True) else None
        self.first = True
        self.f.write("[")

    def write_rows(self, rows):
        for row in rows:
            record = {c: to_json(v) for c, v in zip(self.columns, row)}
            text = json.dumps(record, ensure_ascii=False, indent=self.indent)
            if self.indent:
                text = "\n".join("  " + line for line in text.splitlines())
            self.f.write(("\n" if self.first else ",\n") + text)
            self.first = False

    def close(self):
        self.f.write("\n]\n" if not self.first else "]\n")
        self.f.close()


class XmlWriter(_Writer):
    """Same layout as Navicat: <RECORDS><RECORD><column>value</column>...</RECORD></RECORDS>."""

    def __init__(self, path, columns, options, name):
        super().__init__(path, columns, options, name)
        self.f = open(path, "w", encoding="utf-8")
        self.tags = [re.sub(r"[^\w.-]", "_", c) or "column" for c in columns]
        self.tags = ["_" + t if t[0].isdigit() or t[0] in ".-" else t for t in self.tags]
        self.f.write('<?xml version="1.0" encoding="UTF-8"?>\n<RECORDS>\n')

    def write_rows(self, rows):
        out = []
        for row in rows:
            out.append("<RECORD>\n")
            for tag, v in zip(self.tags, row):
                if v is None:
                    out.append(f"  <{tag}/>\n")
                else:
                    out.append(f"  <{tag}>{xml_escape(to_text(v))}</{tag}>\n")
            out.append("</RECORD>\n")
        self.f.write("".join(out))

    def close(self):
        self.f.write("</RECORDS>\n")
        self.f.close()


class HtmlWriter(_Writer):
    def __init__(self, path, columns, options, name):
        super().__init__(path, columns, options, name)
        self.f = open(path, "w", encoding="utf-8")
        title = html.escape(name)
        self.f.write(
            "<!DOCTYPE html>\n<html>\n<head>\n<meta charset=\"utf-8\">\n"
            f"<title>{title}</title>\n<style>\n"
            "body{font-family:sans-serif;font-size:13px;color:#262626}\n"
            "table{border-collapse:collapse}\n"
            "th,td{border:1px solid #dcdcdc;padding:4px 8px;text-align:left;vertical-align:top}\n"
            "th{background:#f5f5f5}\n.null{color:#999;font-style:italic}\n"
            "</style>\n</head>\n<body>\n"
            f"<h3>{title}</h3>\n<table>\n<tr>"
            + "".join(f"<th>{html.escape(c)}</th>" for c in columns) + "</tr>\n"
        )

    def write_rows(self, rows):
        out = []
        for row in rows:
            cells = ['<td class="null">NULL</td>' if v is None else f"<td>{html.escape(to_text(v))}</td>"
                     for v in row]
            out.append("<tr>" + "".join(cells) + "</tr>\n")
        self.f.write("".join(out))

    def close(self):
        self.f.write("</table>\n</body>\n</html>\n")
        self.f.close()


class XlsxWriter(_Writer):
    MAX_ROWS = 1_048_575  # Excel sheet limit, minus the header row

    def __init__(self, path, columns, options, name):
        super().__init__(path, columns, options, name)
        self.wb = openpyxl.Workbook(write_only=True)
        self.ws = self.wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet1")
        self.count = 0
        if options.get("header", True):
            self.ws.append(columns)

    @staticmethod
    def _cell(v):
        if isinstance(v, (bytes, bytearray)):
            return base64.b64encode(bytes(v)).decode("ascii")
        if isinstance(v, datetime.timedelta):
            return str(v)
        if isinstance(v, decimal.Decimal):
            return float(v) if abs(v) < 1e15 else format(v, "f")
        if isinstance(v, str) and len(v) > 32767:
            return v[:32767]
        return v

    def write_rows(self, rows):
        for row in rows:
            if self.count >= self.MAX_ROWS:
                raise ValueError("Excel supports at most 1,048,576 rows per sheet.")
            self.ws.append([self._cell(v) for v in row])
            self.count += 1

    def close(self):
        self.wb.save(self.path)


class SqlWriter(_Writer):
    """INSERT statements, optionally preceded by DROP/CREATE TABLE."""

    def __init__(self, path, columns, options, name, conn=None, create_sql=None):
        super().__init__(path, columns, options, name)
        self.conn = conn
        self.f = open(path, "w", encoding="utf-8")
        self.per_insert = max(1, int(options.get("rows_per_insert", 100)))
        self.table = q(options.get("target_table") or name)
        self.f.write(_sql_header(options.get("database")))
        if create_sql is not None:
            write_table_ddl(self.f, name, create_sql, options.get("drop", True))
        self.insert = f"INSERT INTO {self.table} ({', '.join(q(c) for c in columns)}) VALUES\n"

    def write_rows(self, rows):
        literal = self.conn.literal
        for i in range(0, len(rows), self.per_insert):
            chunk = rows[i:i + self.per_insert]
            values = ",\n".join("(" + ", ".join(literal(v) for v in row) + ")" for row in chunk)
            self.f.write(self.insert + values + ";\n")

    def close(self):
        self.f.write(_sql_footer())
        self.f.close()


WRITERS = {"csv": DelimitedWriter, "txt": DelimitedWriter, "json": JsonWriter, "xml": XmlWriter,
           "html": HtmlWriter, "xlsx": XlsxWriter, "sql": SqlWriter}


def _sql_header(database=None):
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    head = (f"-- {APP_NAME} {APP_VERSION} SQL dump\n-- Date: {now}\n"
            + (f"-- Database: {database}\n" if database else "") + "\n"
            "SET NAMES utf8mb4;\nSET FOREIGN_KEY_CHECKS = 0;\n\n")
    return head


def _sql_footer():
    return "\nSET FOREIGN_KEY_CHECKS = 1;\n"


def write_table_ddl(f, table, create_sql, drop=True, kind="TABLE"):
    f.write(f"-- ----------------------------\n-- {kind.title()} structure for {table}\n-- ----------------------------\n")
    if drop:
        f.write(f"DROP {kind} IF EXISTS {q(table)};\n")
    f.write(create_sql.rstrip(";") + ";\n\n")


# --- export ------------------------------------------------------------------------
def default_export_options(fmt):
    base = {"header": True, "null_text": "", "encoding": "utf-8"}
    if fmt == "csv":
        base.update(delimiter=",", quote='"', line_ending="\r\n")
    elif fmt == "txt":
        base.update(delimiter="\t", quote="", line_ending="\n")
    elif fmt == "sql":
        base.update(create=True, drop=True, rows_per_insert=100)
    elif fmt == "json":
        base.update(pretty=True)
    return base


def export_query(conn, sql, fmt, path, options, name, progress=_noop, cancelled=_never, total=None,
                 create_sql=None):
    """Stream the rows of `sql` into a file. Returns the number of rows written."""
    conn.ping(reconnect=True)
    with conn.cursor(pymysql.cursors.SSCursor) as cur:
        cur.execute(sql)
        columns = [d[0] for d in cur.description]
        writer_cls = WRITERS[fmt]
        if fmt == "sql":
            writer = writer_cls(path, columns, options, name, conn=conn, create_sql=create_sql)
        else:
            writer = writer_cls(path, columns, options, name)
        done = 0
        try:
            while True:
                if cancelled():
                    raise Cancelled()
                rows = cur.fetchmany(CHUNK)
                if not rows:
                    break
                writer.write_rows(rows)
                done += len(rows)
                progress(done, total, f"{name}: {done:,} rows")
        finally:
            writer.close()  # closing the SSCursor (end of `with`) drains any unread rows
    return done


def export_rows(columns, rows, fmt, path, options, name, conn=None):
    """Export rows that are already in memory (e.g. a query result)."""
    if fmt == "sql":
        writer = SqlWriter(path, columns, options, name, conn=conn)
    else:
        writer = WRITERS[fmt](path, columns, options, name)
    try:
        writer.write_rows([tuple(r) for r in rows])
    finally:
        writer.close()
    return len(rows)


def export_table(conn, database, table, fmt, path, options, progress=_noop, cancelled=_never):
    fqn = f"{q(database)}.{q(table)}"
    total = db.execute(conn, f"SELECT COUNT(*) FROM {fqn}").rows[0][0]
    create_sql = None
    if fmt == "sql" and options.get("create", True):
        create_sql = db.execute(conn, f"SHOW CREATE TABLE {fqn}").rows[0][1]
    opts = dict(options, database=database)
    return export_query(conn, f"SELECT * FROM {fqn}", fmt, path, opts, table, progress, cancelled, total, create_sql)


# --- SQL dump of a whole database ----------------------------------------------------
def dump_database(conn, database, path, with_data=True, progress=_noop, cancelled=_never):
    """Write structure (tables, views, functions, procedures) and optionally data to one .sql file."""
    tables = db.list_tables(conn, database)
    routines = db.list_routines(conn, database)
    base = [t for t, kind in tables if kind != "VIEW"]
    views = [t for t, kind in tables if kind == "VIEW"]
    steps = len(tables) + len(routines)
    done = 0
    literal = conn.literal
    with open(path, "w", encoding="utf-8") as f:
        f.write(_sql_header(database))
        for table in base:
            if cancelled():
                raise Cancelled()
            fqn = f"{q(database)}.{q(table)}"
            write_table_ddl(f, table, db.execute(conn, f"SHOW CREATE TABLE {fqn}").rows[0][1])
            if with_data:
                conn.ping(reconnect=True)
                with conn.cursor(pymysql.cursors.SSCursor) as cur:
                    cur.execute(f"SELECT * FROM {fqn}")
                    columns = ", ".join(q(d[0]) for d in cur.description)
                    rows_written = 0
                    while True:
                        rows = cur.fetchmany(100)
                        if not rows:
                            break
                        if rows_written == 0:
                            f.write(f"-- ----------------------------\n-- Records of {table}\n"
                                    "-- ----------------------------\n")
                        values = ",\n".join("(" + ", ".join(literal(v) for v in row) + ")" for row in rows)
                        f.write(f"INSERT INTO {q(table)} ({columns}) VALUES\n{values};\n")
                        rows_written += len(rows)
                        if cancelled():
                            raise Cancelled()
                    if rows_written:
                        f.write("\n")
            done += 1
            progress(done, steps, f"Table {table}")
        for view in views:
            ddl = db.execute(conn, f"SHOW CREATE VIEW {q(database)}.{q(view)}").rows[0][1]
            write_table_ddl(f, view, ddl, kind="VIEW")
            done += 1
            progress(done, steps, f"View {view}")
        for name, rtype in routines:
            ddl = db.routine_definition(conn, database, name, rtype)
            f.write(f"-- ----------------------------\n-- {rtype.title()} structure for {name}\n"
                    f"-- ----------------------------\nDROP {rtype} IF EXISTS {q(name)};\n"
                    f"DELIMITER ;;\n{ddl.rstrip(';')}\n;;\nDELIMITER ;\n\n")
            done += 1
            progress(done, steps, f"{rtype.title()} {name}")
        f.write(_sql_footer())
    return done


def execute_sql_file(conn, path, continue_on_error=False, progress=_noop, cancelled=_never, log=_noop):
    """Run every statement in an .sql file. Returns (ok_count, error_count)."""
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        statements = db.split_statements(f.read())
    ok = errors = 0
    for i, stmt in enumerate(statements, 1):
        if cancelled():
            raise Cancelled()
        try:
            db.execute_all(conn, stmt)
            ok += 1
        except Exception as e:  # noqa: BLE001
            errors += 1
            log(f"[ERR] {e}\n      {' '.join(stmt.split())[:200]}")
            if not continue_on_error:
                raise
        if i % 20 == 0 or i == len(statements):
            progress(i, len(statements), f"{i:,} of {len(statements):,} statements")
    return ok, errors


# --- readers (import sources) ----------------------------------------------------------
def read_source(fmt, path, options):
    """Return (columns, row iterator) for an import source. Rows are lists of str/None values."""
    if fmt in ("csv", "txt"):
        return _read_delimited(path, options)
    if fmt == "json":
        return _read_json(path)
    if fmt == "xlsx":
        return _read_xlsx(path, options)
    raise ValueError(f"Unsupported format: {fmt}")


def _read_delimited(path, options):
    encoding = options.get("encoding", "utf-8-sig")
    delimiter = options.get("delimiter", ",")
    quote = options.get("quote", '"')
    skip = int(options.get("skip", 0))
    header = options.get("header", True)

    def rows():
        with open(path, newline="", encoding=encoding, errors="replace") as f:
            reader = csv.reader(f, delimiter=delimiter, quotechar=quote or '"',
                                quoting=csv.QUOTE_MINIMAL if quote else csv.QUOTE_NONE)
            for _ in range(skip + (1 if header else 0)):
                next(reader, None)
            for row in reader:
                if row:
                    yield row

    with open(path, newline="", encoding=encoding, errors="replace") as f:
        reader = csv.reader(f, delimiter=delimiter, quotechar=quote or '"',
                            quoting=csv.QUOTE_MINIMAL if quote else csv.QUOTE_NONE)
        for _ in range(skip):
            next(reader, None)
        first = next(reader, [])
    columns = [c.strip() or f"column{i + 1}" for i, c in enumerate(first)] if header else \
        [f"column{i + 1}" for i in range(len(first))]
    return _dedupe(columns), rows()


def _read_json(path):
    with open(path, encoding="utf-8-sig") as f:
        data = json.load(f)
    if isinstance(data, dict):
        # {"RECORDS": [...]} (Navicat) or any single key holding a list
        lists = [v for v in data.values() if isinstance(v, list)]
        data = lists[0] if len(lists) == 1 else [data]
    if not isinstance(data, list):
        raise ValueError("The JSON file must contain an array of objects.")
    columns = []
    for rec in data[:1000]:
        if isinstance(rec, dict):
            for k in rec:
                if k not in columns:
                    columns.append(k)
    if not columns:
        raise ValueError("No objects found in the JSON array.")

    def cell(v):
        if v is None:
            return None
        if isinstance(v, (dict, list)):
            return json.dumps(v, ensure_ascii=False)
        if isinstance(v, bool):
            return "1" if v else "0"
        return str(v)

    def rows():
        for rec in data:
            if isinstance(rec, dict):
                yield [cell(rec.get(c)) for c in columns]

    return columns, rows()


def xlsx_sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def _read_xlsx(path, options):
    header = options.get("header", True)
    skip = int(options.get("skip", 0))

    def cell(v):
        if v is None:
            return None
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)

    def open_sheet():
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[options["sheet"]] if options.get("sheet") else wb.active
        return wb, ws.iter_rows(values_only=True)

    wb, it = open_sheet()
    try:
        for _ in range(skip):
            next(it, None)
        first = next(it, ()) or ()
    finally:
        wb.close()
    width = len(first)
    columns = [str(c).strip() if c is not None else f"column{i + 1}" for i, c in enumerate(first)] if header \
        else [f"column{i + 1}" for i in range(width)]

    def rows():
        wb2, it2 = open_sheet()
        try:
            for _ in range(skip + (1 if header else 0)):
                next(it2, None)
            for row in it2:
                if row and any(v is not None for v in row):
                    # openpyxl drops empty trailing cells; keep every row as wide as the header
                    yield ([cell(v) for v in row] + [None] * width)[:width]
        finally:
            wb2.close()

    return _dedupe(columns), rows()


def _dedupe(columns):
    seen, out = {}, []
    for c in columns:
        n = seen.get(c, 0)
        seen[c] = n + 1
        out.append(c if n == 0 else f"{c}_{n + 1}")
    return out


def preview(fmt, path, options, limit=50):
    columns, rows = read_source(fmt, path, options)
    sample = []
    for row in rows:
        sample.append(row)
        if len(sample) >= limit:
            break
    return columns, sample


# --- import ----------------------------------------------------------------------------
_INT = re.compile(r"^[+-]?\d+$")
_DEC = re.compile(r"^[+-]?\d*\.\d+$|^[+-]?\d+\.\d*$")
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2}(\.\d+)?)?$")


def infer_type(values):
    """Pick a MySQL column type that fits all sample values."""
    vals = [v for v in values if v not in (None, "")]
    if not vals:
        return "VARCHAR(255)"
    if all(_INT.match(v) for v in vals):
        big = max(abs(int(v)) for v in vals)
        return "BIGINT" if big > 2_147_483_647 else "INT"
    if all(_INT.match(v) or _DEC.match(v) for v in vals):
        scale = max((len(v.split(".")[1]) if "." in v else 0) for v in vals)
        whole = max(len(v.lstrip("+-").split(".")[0]) for v in vals)
        scale = min(scale, 10)
        return f"DECIMAL({min(whole + scale + 2, 65)},{scale})"
    if all(_DATE.match(v) for v in vals):
        return "DATE"
    if all(_DATETIME.match(v) for v in vals):
        return "DATETIME"
    longest = max(len(v) for v in vals)
    if longest <= 255:
        return "VARCHAR(255)"
    return "TEXT" if longest <= 65535 else "LONGTEXT"


def create_table_sql(database, table, columns, types, keys=()):
    defs = [f"{q(c)} {t} {'NOT NULL' if c in keys else 'NULL'}" for c, t in zip(columns, types)]
    if keys:
        defs.append(f"PRIMARY KEY ({', '.join(q(k) for k in keys)})")
    body = ",\n  ".join(defs)
    return (f"CREATE TABLE {q(database)}.{q(table)} (\n  {body}\n) "
            "ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci")


def import_rows(conn, database, table, mapping, rows, mode="append", key_columns=(), empty_as_null=True,
                continue_on_error=False, use_transaction=True, progress=_noop, cancelled=_never, log=_noop,
                total=None):
    """Import rows into database.table.

    mapping: list of (source index, target column). key_columns: target columns that identify a
    record (needed by update / append_update / delete). Returns (processed, errors).
    """
    if not mapping:
        raise ValueError("No fields are mapped.")
    fqn = f"{q(database)}.{q(table)}"
    targets = [t for _, t in mapping]
    keys = [k for k in key_columns if k in targets]
    if mode in ("update", "delete", "append_update") and not keys:
        raise ValueError("This import mode needs the table's primary key columns to be mapped.")
    cols = ", ".join(q(t) for t in targets)
    marks = ", ".join(["%s"] * len(targets))
    others = [t for t in targets if t not in keys]
    key_idx = [targets.index(k) for k in keys]
    other_idx = [targets.index(t) for t in others]
    where = " AND ".join(f"{q(k)} = %s" for k in keys)

    if mode in ("append", "copy"):
        sql, shape = f"INSERT INTO {fqn} ({cols}) VALUES ({marks})", None
    elif mode == "append_update":
        updates = ", ".join(f"{q(t)} = VALUES({q(t)})" for t in others) or f"{q(keys[0])} = {q(keys[0])}"
        sql, shape = f"INSERT INTO {fqn} ({cols}) VALUES ({marks}) ON DUPLICATE KEY UPDATE {updates}", None
    elif mode == "update":
        if not others:
            raise ValueError("Map at least one non-key field to update.")
        sets = ", ".join(f"{q(t)} = %s" for t in others)
        sql, shape = f"UPDATE {fqn} SET {sets} WHERE {where}", other_idx + key_idx
    elif mode == "delete":
        sql, shape = f"DELETE FROM {fqn} WHERE {where}", key_idx
    else:
        raise ValueError(f"Unknown import mode: {mode}")

    def convert(row):
        values = []
        for src, _ in mapping:
            v = row[src] if src < len(row) else None
            if empty_as_null and v == "":
                v = None
            values.append(v)
        return values if shape is None else [values[i] for i in shape]

    conn.ping(reconnect=True)
    processed = errors = 0
    if use_transaction:
        conn.begin()
    try:
        if mode == "copy":
            db.execute(conn, f"DELETE FROM {fqn}")
            log(f"[MSG] Deleted existing records from {table}")
        batch = []

        def flush():
            nonlocal processed, errors
            if not batch:
                return
            try:
                with conn.cursor() as cur:
                    cur.executemany(sql, batch)
            except Exception as e:  # noqa: BLE001
                if not continue_on_error:
                    raise RuntimeError(f"Error near record {processed + 1}: {e}") from e
                # Retry one by one so only the bad records are skipped.
                for i, params in enumerate(batch):
                    try:
                        with conn.cursor() as cur:
                            cur.execute(sql, params)
                    except Exception as row_err:  # noqa: BLE001
                        errors += 1
                        log(f"[ERR] Record {processed + i + 1}: {row_err}")
            processed += len(batch)
            batch.clear()
            progress(processed, total, f"{processed:,} records processed")

        for row in rows:
            if cancelled():
                raise Cancelled()
            batch.append(convert(row))
            if len(batch) >= 500:
                flush()
        flush()
        if use_transaction:
            conn.commit()
    except BaseException:
        if use_transaction:
            conn.rollback()
            log("[MSG] Rolled back: no records were imported")
        raise
    return processed, errors


def default_path(directory, name, fmt):
    return os.path.join(directory, f"{name}.{EXPORT_FORMATS[fmt][1]}")
